#!/usr/bin/env python3
"""
Pembuat gambar & tabel artikel yang tidak dihasilkan evaluate_accuracy.py.

Subperintah:
  mtcnn  → Gambar 2  : grid foto asli (dengan kotak deteksi) vs crop MTCNN 160×160
  pose   → Gambar 6  : detection rate MTCNN per pose enrollment
                       (membaca detection_log.csv keluaran evaluate_accuracy.py)
  sus    → Gambar 11 + Tabel 6 : skor SUS per item dari kuesioner_sus.xlsx

Pemakaian:
  python scripts/buat_gambar.py mtcnn --images a.jpg b.jpg c.jpg d.jpg --out ./gambar/gambar2.png
  python scripts/buat_gambar.py pose  --log ./eval_640/detection_log.csv --out ./gambar/gambar6.png
  python scripts/buat_gambar.py sus   --xlsx ./pengujian/kuesioner_sus.xlsx --out ./gambar

Semua gambar diekspor PNG 300 dpi.
"""

import argparse
import csv
import os
import sys

DETECTOR_BACKEND = os.getenv("DETECTOR_BACKEND", "mtcnn")
POSES = ["depan", "kiri", "kanan", "menunduk", "mendongak"]


def _plt():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    return plt


def _save(fig, out_path, rect=(0, 0, 1, 1)):
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    fig.tight_layout(rect=rect)
    fig.savefig(out_path, dpi=300)
    print(f"  ✓ {out_path}")


def _write_csv(path, rows):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    print(f"  ✓ {path}")


# ── Gambar 2: preprocessing MTCNN ────────────────────────────────────────────
def cmd_mtcnn(args):
    import cv2
    from deepface import DeepFace

    plt = _plt()
    n = len(args.images)
    fig, axes = plt.subplots(2, n, figsize=(2.6 * n, 5.6), squeeze=False)

    for i, path in enumerate(args.images):
        img = cv2.imread(path)
        if img is None:
            sys.exit(f"Gagal membaca gambar: {path}")
        faces = DeepFace.extract_faces(img_path=path, detector_backend=DETECTOR_BACKEND,
                                       enforce_detection=True, align=True)
        face = faces[0]
        area = face["facial_area"]

        # Baris atas: foto asli + kotak deteksi
        shown = img.copy()
        x, y, w, h = area["x"], area["y"], area["w"], area["h"]
        cv2.rectangle(shown, (x, y), (x + w, y + h), (0, 255, 0), max(2, img.shape[1] // 200))
        # Baris bawah: wajah ter-crop & ter-align, dinormalisasi ke 160×160 (input Facenet512)
        crop = cv2.resize((face["face"] * 255).astype("uint8"), (160, 160),
                          interpolation=cv2.INTER_AREA)
        crop = cv2.cvtColor(crop, cv2.COLOR_RGB2BGR)

        if args.blur:
            # Untuk publikasi tanpa izin wajah: samarkan kedua baris.
            k = max(31, (min(shown.shape[:2]) // 8) | 1)
            shown = cv2.GaussianBlur(shown, (k, k), 0)
            crop = cv2.GaussianBlur(crop, (31, 31), 0)

        axes[0][i].imshow(cv2.cvtColor(shown, cv2.COLOR_BGR2RGB))
        axes[0][i].set_title(args.labels[i] if args.labels else os.path.basename(path), fontsize=9)
        axes[1][i].imshow(cv2.cvtColor(crop, cv2.COLOR_BGR2RGB))
        axes[1][i].set_title("160×160", fontsize=9)
        for r in (0, 1):
            axes[r][i].axis("off")

    axes[0][0].text(-0.08, 0.5, "Asli + deteksi", transform=axes[0][0].transAxes,
                    rotation=90, va="center", ha="right", fontsize=10)
    axes[1][0].text(-0.08, 0.5, "Crop MTCNN", transform=axes[1][0].transAxes,
                    rotation=90, va="center", ha="right", fontsize=10)
    _save(fig, args.out)


# ── Gambar 6: detection rate per pose ────────────────────────────────────────
def cmd_pose(args):
    with open(args.log, newline="") as f:
        rows = [r for r in csv.DictReader(f) if r["split"] == "enroll"]
    if not rows:
        sys.exit(f"Tidak ada baris enroll di {args.log}")

    counts = {}
    for r in rows:
        pose = r["pose"] if r["pose"] in POSES else "lainnya"
        ok, total = counts.get(pose, (0, 0))
        counts[pose] = (ok + int(r["terdeteksi"]), total + 1)
    if "lainnya" in counts:
        print(f"  [WARN] {counts['lainnya'][1]} foto enroll tidak bernama pose "
              f"({', '.join(POSES)}) → dikelompokkan 'lainnya'")

    order = [p for p in POSES + ["lainnya"] if p in counts]
    summary = []
    for p in order:
        ok, total = counts[p]
        summary.append({"pose": p, "total_foto": total, "terdeteksi": ok,
                        "gagal": total - ok, "detection_rate_percent": round(100 * ok / total, 2)})
    ok_all = sum(s["terdeteksi"] for s in summary)
    n_all = sum(s["total_foto"] for s in summary)
    summary.append({"pose": "TOTAL", "total_foto": n_all, "terdeteksi": ok_all,
                    "gagal": n_all - ok_all, "detection_rate_percent": round(100 * ok_all / n_all, 2)})
    _write_csv(os.path.splitext(args.out)[0] + ".csv", summary)

    plt = _plt()
    fig, ax = plt.subplots(figsize=(7, 4.5))
    labels = [s["pose"].capitalize() for s in summary[:-1]]
    ok = [s["terdeteksi"] for s in summary[:-1]]
    fail = [s["gagal"] for s in summary[:-1]]
    ax.bar(labels, ok, color="#2a9d8f", label="Terdeteksi")
    ax.bar(labels, fail, bottom=ok, color="#e76f51", label="Gagal terdeteksi")
    for i, s in enumerate(summary[:-1]):
        ax.text(i, s["total_foto"], f"{s['detection_rate_percent']:.1f}%",
                ha="center", va="bottom", fontsize=9)
    ax.set_ylabel("Jumlah foto")
    ax.set_title(f"Detection Rate MTCNN per Pose (total {summary[-1]['detection_rate_percent']:.1f}%)")
    ax.set_ylim(0, max(s["total_foto"] for s in summary[:-1]) * 1.15)
    ax.legend(loc="lower right")
    ax.grid(axis="y", alpha=0.3)
    _save(fig, args.out)


# ── Gambar 11 + Tabel 6: SUS ─────────────────────────────────────────────────
def sus_category(score):
    if score >= 80.3:
        return "Excellent (A)"
    if score >= 68:
        return "Good (B)"
    if score >= 51:
        return "OK (C)"
    return "Poor"


def cmd_sus(args):
    import openpyxl

    wb = openpyxl.load_workbook(args.xlsx)
    data, rekap = wb["Data"], wb["Rekap Item"]
    statements = [rekap.cell(row=3 + i, column=2).value for i in range(10)]

    # Hitung ulang dari jawaban mentah (rumus Excel tidak dievaluasi oleh openpyxl).
    responses = []
    for row in data.iter_rows(min_row=4, values_only=True):
        label, answers = row[0], row[1:11]
        if label == "CONTOH" or not all(isinstance(a, (int, float)) for a in answers):
            continue
        if not all(1 <= a <= 5 for a in answers):
            sys.exit(f"Responden {label}: jawaban di luar skala 1–5: {answers}")
        responses.append([int(a) for a in answers])
    if not responses:
        sys.exit("Belum ada responden lengkap (P1–P10) di sheet Data.")

    def contribution(item_idx, score):
        # Item ganjil (positif): skor−1; item genap (negatif): 5−skor.
        return score - 1 if item_idx % 2 == 0 else 5 - score

    scores = [sum(contribution(i, a) for i, a in enumerate(r)) * 2.5 for r in responses]
    mean_score = sum(scores) / len(scores)

    table = []
    for i in range(10):
        avg = sum(r[i] for r in responses) / len(responses)
        table.append({"no": i + 1, "pernyataan": statements[i],
                      "skor_rerata": round(avg, 2), "bobot_kontribusi": round(contribution(i, avg), 2)})
    # Total dari skor tak-dibulatkan → sama persis dengan rata-rata skor SUS responden.
    table.append({"no": "", "pernyataan": "Total bobot / Skor SUS (×2,5)",
                  "skor_rerata": "", "bobot_kontribusi": round(mean_score, 2)})

    os.makedirs(args.out, exist_ok=True)
    _write_csv(os.path.join(args.out, "tabel6_sus_per_item.csv"), table)

    plt = _plt()
    fig, ax = plt.subplots(figsize=(8, 5.5))
    items = table[:10]
    # Warna berdasarkan BOBOT KONTRIBUSI (0–4), bukan skor mentah: pada item genap
    # (negatif) skor tinggi justru buruk. Ambang 3 & 2 setara skor 4 & 3 di item positif.
    colors = ["#2a9d8f" if t["bobot_kontribusi"] >= 3 else
              "#e9c46a" if t["bobot_kontribusi"] >= 2 else "#e76f51" for t in items]
    labels = [f"P{t['no']}{' (−)' if t['no'] % 2 == 0 else ''}" for t in items]
    y = list(range(10))[::-1]
    ax.barh(y, [t["skor_rerata"] for t in items], color=colors)
    for yi, t in zip(y, items):
        ax.text(t["skor_rerata"] + 0.05, yi, f"{t['skor_rerata']:.2f}", va="center", fontsize=9)
    ax.set_yticks(y)
    ax.set_yticklabels(labels)
    ax.set_xlim(1, 5.4)
    ax.set_xlabel("Skor rerata (1–5)")
    ax.set_title(f"Skor SUS per Item — total {mean_score:.2f} ({sus_category(mean_score)}), "
                 f"n = {len(responses)}")
    ax.grid(axis="x", alpha=0.3)
    fig.text(0.01, 0.01, "(−) = pernyataan negatif; skor rendah = baik. "
             "Warna: hijau baik, kuning sedang, merah perlu perbaikan.", fontsize=8, color="0.4")
    _save(fig, os.path.join(args.out, "gambar11_sus.png"), rect=(0, 0.04, 1, 1))

    print("\n" + "=" * 52)
    print(f"  Responden      : {len(responses)}")
    print(f"  Skor SUS rerata: {mean_score:.2f}  →  {sus_category(mean_score)}")
    print("=" * 52)


def main():
    ap = argparse.ArgumentParser(description="Pembuat gambar & tabel artikel E-Voting")
    sub = ap.add_subparsers(dest="cmd", required=True)

    m = sub.add_parser("mtcnn", help="Gambar 2: foto asli vs crop MTCNN 160×160")
    m.add_argument("--images", nargs="+", required=True, help="Beberapa foto contoh (disarankan 4–6)")
    m.add_argument("--labels", nargs="+", help="Judul kolom (mis. kondisi), sejumlah --images")
    m.add_argument("--blur", action="store_true", help="Samarkan wajah (bila tanpa izin publikasi)")
    m.add_argument("--out", default="./gambar/gambar2_mtcnn.png")

    p = sub.add_parser("pose", help="Gambar 6: detection rate per pose")
    p.add_argument("--log", required=True, help="detection_log.csv dari evaluate_accuracy.py")
    p.add_argument("--out", default="./gambar/gambar6_pose.png")

    s = sub.add_parser("sus", help="Gambar 11 + Tabel 6: skor SUS per item")
    s.add_argument("--xlsx", default="./pengujian/kuesioner_sus.xlsx")
    s.add_argument("--out", default="./gambar", help="Folder output")

    args = ap.parse_args()
    if args.cmd == "mtcnn" and args.labels and len(args.labels) != len(args.images):
        sys.exit("--labels harus sejumlah --images")
    {"mtcnn": cmd_mtcnn, "pose": cmd_pose, "sus": cmd_sus}[args.cmd](args)


if __name__ == "__main__":
    main()
